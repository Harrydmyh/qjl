#!/usr/bin/env python3
"""帮卖广场「累计订单 1000+」活动报表 v4。

v4 变更（用户纠正链接格式）：
  链接改用微信群接龙小程序短链文本格式：
      #小程序://群接龙/{16位短码}
  由平台短链接口按活动逐一生成（不是从 actId 推的，必须真调接口）：
      POST apipro/wxa-link/pc/gen_short_link
      body（平铺，不要 param 包裹）：
        {pageTitle: "", pageUrl: "pro/pages/seq-detail/detail-group-buy/detail-group-buy?actId=<actId>&fromShare=1&shareByLink=1",
         appId: "wx059cd327295ab444", isPermanent: true}
  pageUrl 路径来自网页端 getMiniProgramPageUrl() 的同款构造；appId 生产环境值
  wx059cd327295ab444（即群接龙小程序）。生成的短链带本地缓存，重跑不重复造。

沿用 v2/v3：按生效时间逐月扫窗（2019-01 → 今）覆盖老团；窗内热度早停；
已暂停/已下架活动接口不返回数据，天然进不了报表；「最近帮卖日期」来自平台帮卖子团记录。

用法：/opt/homebrew/bin/python3 hot_acts_report.py [--min-orders 1000]
"""
import argparse
import json
import os
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta

import core
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ZW_CHARS = ("\u200b", "\u200c", "\u200d", "\ufeff")
APP_ID = "wx059cd327295ab444"
SHORT_LINK_API = "https://apipro.qunjielong.com/wxa-link/pc/gen_short_link"
LINK_CACHE = "data/short_link_cache.json"


def clean(s):
    s = s or ""
    for z in ZW_CHARS:
        s = s.replace(z, "")
    return s.strip()


def initial_date_from_act_id(act_id, effect_ms=0):
    aid = str(act_id or "")
    if len(aid) >= 8 and aid[:6].isdigit():
        y, m, d = 2000 + int(aid[:2]), int(aid[2:4]), int(aid[4:6])
        try:
            dt = datetime(y, m, d)
            if dt <= datetime.now() + timedelta(days=1):
                return dt.strftime("%Y-%m-%d")
        except ValueError:
            pass
    if effect_ms:
        return datetime.fromtimestamp(effect_ms / 1000).strftime("%Y-%m-%d")
    return ""


def month_windows(start, end):
    wins, cur = [], datetime(start.year, start.month, 1)
    while cur <= end:
        nxt = datetime(cur.year + (cur.month == 12), (cur.month % 12) + 1, 1)
        wins.append((cur, nxt + timedelta(days=1)))
        cur = nxt
    return wins


def pull_window(w0, w1, gh_id, headers, min_orders, stop_streak, page_cap):
    session = core.make_session()
    start_ts = int(w0.timestamp() * 1000)
    end_ts = int(w1.timestamp() * 1000)
    hits, listed, detailed, page, streak = {}, 0, 0, 1, 0
    while page <= page_cap:
        try:
            _, items, has_next = core._fetch_one_page(
                session, headers, gh_id, page, [], [2, 3], start_ts, end_ts)
        except Exception:
            break
        if not items:
            break
        page_ids = [it["actId"] for it in items if it.get("actId")]
        listed += len(page_ids)
        det = core._fetch_details_batch(session, headers, gh_id, page_ids) if page_ids else []
        detailed += len(det)
        page_max = 0
        for d in det:
            orders = (d.get("actStat") or {}).get("totalOrderNum", 0) or 0
            page_max = max(page_max, orders)
            if orders < min_orders:
                continue
            if d.get("activityStatus") != 10 or d.get("pauseSw", 5) != 5:
                continue
            aid = str(d.get("actId") or "")
            if not aid or aid in hits:
                continue
            goods = (d.get("basicFeedGoodDTOList") or [{}])[0]
            hits[aid] = {
                "actId": aid,
                "name": clean(goods.get("goodsName")) or clean(d.get("actName")),
                "supplier": d.get("ghName") or "",
                "orders": orders,
                "initialDate": initial_date_from_act_id(aid, d.get("effectTime")),
            }
        if page_max < min_orders:
            streak += 1
            if streak >= stop_streak:
                break
        else:
            streak = 0
        if not has_next:
            break
        page += 1
    return hits, listed, detailed


def gen_short_link(aid, headers, tries=2):
    """按活动生成微信群接龙小程序短链；失败返回 None（调用方用路径格式兜底）。"""
    page_url = (f"pro/pages/seq-detail/detail-group-buy/detail-group-buy"
                f"?actId={aid}&fromShare=1&shareByLink=1")
    body = {"pageTitle": "", "pageUrl": page_url, "appId": APP_ID, "isPermanent": True}
    for i in range(tries):
        try:
            r = core.make_session().post(SHORT_LINK_API, json=body, headers=headers, timeout=20)
            d = (r.json().get("data") or {})
            link = d.get("link") or ""
            if link and not d.get("errCode"):
                return link
        except Exception:
            pass
        time.sleep(1.0 * (i + 1))
    return None


def build_links(hit_ids, headers, workers=6):
    """带缓存地批量生成短链 → {actId: link}"""
    cache = {}
    if os.path.exists(LINK_CACHE):
        try:
            cache = json.load(open(LINK_CACHE, encoding="utf-8"))
        except Exception:
            cache = {}
    todo = [a for a in hit_ids if a not in cache]
    print(f"短链生成: 缓存命中 {len(hit_ids) - len(todo)}，新生成 {len(todo)}", flush=True)

    def one(aid):
        return aid, gen_short_link(aid, headers)

    ok = fail = 0
    if todo:
        with ThreadPoolExecutor(max_workers=workers) as ex:
            futs = [ex.submit(one, a) for a in todo]
            for fut in as_completed(futs):
                try:
                    aid, link = fut.result()
                except Exception:
                    aid, link = None, None
                if aid and link:
                    cache[aid] = link
                    ok += 1
                else:
                    fail += 1
        try:
            os.makedirs(os.path.dirname(LINK_CACHE), exist_ok=True)
            json.dump(cache, open(LINK_CACHE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        except Exception as e:
            print("缓存写入失败:", e)
    print(f"短链生成完成: 成功 {ok}，失败 {fail}", flush=True)
    return cache, fail


def load_local_help_sale_dates():
    out = {}
    try:
        for e in json.load(open("data/open_log.json", encoding="utf-8")):
            t, aid = e.get("time"), str(e.get("actId") or "")
            if t and aid:
                out.setdefault(aid, []).append(t)
    except Exception:
        pass
    try:
        for run in json.load(open("data/auto_open_log.json", encoding="utf-8")):
            for ok in (run.get("ok") or []):
                t, aid = ok.get("time"), str(ok.get("actId") or "")
                if t and aid:
                    out.setdefault(aid, []).append(t)
    except Exception:
        pass
    return out


def fetch_api_help_sale_dates(hit_ids, gh_id, headers, workers=6):
    def one(aid):
        session = core.make_session()
        latest, page = None, 1
        while page <= 5:
            try:
                r = session.post(
                    "https://apipro.qunjielong.com/activity-accessory/gh_activity/list_gh_child_activities",
                    json={"ghId": gh_id, "pactId": aid, "page": page, "pageSize": 50},
                    headers=headers, timeout=15)
                lst = r.json().get("data") or []
            except Exception:
                break
            for ch in lst:
                ct = ch.get("createTime")
                if ct and (latest is None or ct > latest):
                    latest = ct
            if len(lst) < 50:
                break
            page += 1
        return aid, latest

    result = {}
    with ThreadPoolExecutor(max_workers=workers) as ex:
        futs = [ex.submit(one, aid) for aid in hit_ids]
        for fut in as_completed(futs):
            try:
                aid, latest = fut.result()
                if latest:
                    result[aid] = latest
            except Exception:
                pass
    return result


def merge_help_sale_date(api_ms, local_times):
    cands = []
    if api_ms:
        try:
            cands.append(datetime.fromtimestamp(api_ms / 1000))
        except Exception:
            pass
    for t in (local_times or []):
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
            try:
                cands.append(datetime.strptime(str(t), fmt))
                break
            except ValueError:
                continue
    if not cands:
        return "未帮卖"
    return max(cands).strftime("%Y-%m-%d")


def fallback_link(aid):
    return (f"#小程序：//群接龙?pro/pages/seq-detail/detail-group-buy/"
            f"detail-group-buy?actId={aid}")


def build_xlsx(rows, stats, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "千单以上活动"
    ws.append(["品名", "可点击链接", "供货商", "订单数量", "初始开团日期", "最近帮卖日期"])
    hfill = PatternFill("solid", fgColor="1F4E79")
    hfont = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[1]:
        c.fill, c.font, c.border = hfill, hfont, border
        c.alignment = Alignment(horizontal="center", vertical="center")
    link_font = Font(color="0563C1", underline="single")
    for r in rows:
        ws.append([r["name"], r["link"], r["supplier"], r["orders"],
                   r["initialDate"], r["helpSaleDate"]])
        row = ws.max_row
        ws.cell(row=row, column=2).font = link_font
        ws.cell(row=row, column=4).number_format = "#,##0"
        for col in range(1, 7):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            if col in (4, 5, 6):
                cell.alignment = Alignment(horizontal="center")
    for i, w in enumerate([50, 40, 20, 12, 14, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("说明")
    for line in [
        f"数据来源：群接龙帮卖广场，按生效时间逐月扫描 {stats['since']} 至 {stats['now']}，按累计订单数降序。",
        f"筛选条件：累计订单数（actStat.totalOrderNum）≥ {stats['min_orders']} 单，且活动进行中。",
        f"扫描规模：{stats['windows']} 个月度窗口，列表活动 {stats['listed']} 条，其中 {stats['invisible']} 条接口不返回数据"
        "（已暂停/已下架的活动，不进报表）；最终命中 {n} 条。".format(n=len(rows)),
        "「初始开团日期」取自活动 ID 的创建日期前缀（该活动最初开团的日期，非最近复团日）。",
        "「可点击链接」为按活动逐一生成的微信群接龙小程序短链（#小程序://群接龙/短码，永久有效），"
        "在微信里点开或粘贴即可唤起小程序打开对应活动。",
        "「最近帮卖日期」= 我们帮卖该源活动的最近日期（平台帮卖子团创建时间 + 工具开团记录取最晚）；没有帮卖记录显示「未帮卖」。",
        f"生成时间：{stats['now']}",
    ]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 110
    wb.save(out_path)
    print("saved:", out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--min-orders", type=int, default=1000)
    ap.add_argument("--since", default="2019-01-01")
    ap.add_argument("--stop-streak", type=int, default=6)
    ap.add_argument("--page-cap", type=int, default=40)
    ap.add_argument("--workers", type=int, default=6)
    ap.add_argument("--out", default=None)
    ap.add_argument("--skip-local-log", action="store_true",
                    help="不合并 open_log/auto_open_log 的帮卖记录（换账号时用：那些记录是旧账号干的，不能算到新账号头上）")
    args = ap.parse_args()

    config = json.load(open("data/config.json", encoding="utf-8"))
    headers = core.make_headers(config)
    gh_id = config["gh_id"]
    now = datetime.now()
    since = datetime.strptime(args.since, "%Y-%m-%d")
    windows = month_windows(since, now)
    print(f"扫描 {len(windows)} 个月度窗口（{args.since} → {now:%Y-%m-%d}），阈值 {args.min_orders} 单", flush=True)

    hits = {}
    listed = detailed = 0
    done = 0
    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        futs = [ex.submit(pull_window, w0, w1, gh_id, headers,
                          args.min_orders, args.stop_streak, args.page_cap)
                for w0, w1 in windows]
        for fut in as_completed(futs):
            try:
                wh, wl, wd = fut.result()
            except Exception as e:
                print(f"  窗口失败: {type(e).__name__} {str(e)[:80]}", flush=True)
                continue
            for aid, row in wh.items():
                hits.setdefault(aid, row)
            listed += wl
            detailed += wd
            done += 1
            if done % 24 == 0 or done == len(windows):
                print(f"  进度 {done}/{len(windows)} 窗，累计命中 {len(hits)}", flush=True)

    invisible = listed - detailed
    print(f"\n扫描完成: 列表 {listed} | 可见 {detailed} | 隐身(暂停/已下架) {invisible} | 命中 {len(hits)}", flush=True)

    # 1) 小程序短链（逐活动生成）
    link_cache, link_fail = build_links(list(hits.keys()), headers, args.workers)
    # 2) 最近帮卖日期
    print("拉取平台侧帮卖子团记录...", flush=True)
    api_dates = fetch_api_help_sale_dates(list(hits.keys()), gh_id, headers, args.workers)
    local_dates = {} if args.skip_local_log else load_local_help_sale_dates()
    if args.skip_local_log:
        print("（--skip-local-log：帮卖日期只算当前账号的平台侧记录）", flush=True)

    rows = []
    helped = 0
    for aid, r in hits.items():
        hd = merge_help_sale_date(api_dates.get(aid), local_dates.get(aid))
        if hd != "未帮卖":
            helped += 1
        rows.append({**r, "link": link_cache.get(aid) or fallback_link(aid), "helpSaleDate": hd})
    rows.sort(key=lambda r: -r["orders"])
    print(f"帮卖日期: {helped} 条有记录，{len(rows) - helped} 条未帮卖；链接兜底 {link_fail} 条", flush=True)

    stats = {"since": args.since, "now": now.strftime("%Y-%m-%d %H:%M"),
             "min_orders": args.min_orders, "windows": len(windows),
             "listed": listed, "invisible": invisible}
    out = args.out or f"帮卖广场_千单以上活动_{now.strftime('%Y%m%d_%H%M')}.xlsx"
    build_xlsx(rows, stats, out)
    print("\n===== TOP 10（含链接样例）=====")
    for r in rows[:10]:
        print(f"{r['orders']:>7} | {r['supplier'][:10]:<10} | 帮卖:{r['helpSaleDate']:<10} | {r['link']}")


if __name__ == "__main__":
    main()
